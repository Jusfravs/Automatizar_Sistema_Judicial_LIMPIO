import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from src.gestor_casos import GestorCasos
from src.reparacion_excel import reparar_tablas_vacias


class TablasExcelTests(unittest.TestCase):
    def test_exportador_no_crea_tablas_sin_filas(self):
        with tempfile.TemporaryDirectory() as temporal:
            ruta = Path(temporal) / "resultado.xlsx"
            wb = Workbook()
            reporte = wb.active
            reporte.title = "Reporte"
            encabezados = [f"C{indice}" for indice in range(1, 21)]
            reporte.append(encabezados)
            for fila in (2, 3, 4, 5):
                reporte.cell(fila, 1, "titulo")
            for indice, nombre in enumerate(encabezados, 1):
                reporte.cell(6, indice, nombre)
            reporte.append(list(range(1, 21)))
            for nombre, con_datos in (
                ("PARA CARGA", True),
                ("HISTORIAL HITOS", False),
                ("REVISION IA", False),
            ):
                hoja = wb.create_sheet(nombre)
                hoja.append(["A", "B", "C", "D", "E"])
                if con_datos:
                    hoja.append([1, 2, 3, 4, 5])
            wb.save(ruta)

            GestorCasos._formatear_errores_excel(ruta)
            resultado = load_workbook(ruta)
            self.assertEqual(len(resultado["Reporte"].tables), 1)
            self.assertEqual(len(resultado["PARA CARGA"].tables), 1)
            self.assertEqual(len(resultado["HISTORIAL HITOS"].tables), 0)
            self.assertEqual(len(resultado["REVISION IA"].tables), 0)
            self.assertEqual(resultado["REVISION IA"]["A1"].value, "A")
            resultado.close()

    def test_reparacion_conserva_celdas_y_tablas_con_datos(self):
        wb = Workbook()
        con_datos = wb.active
        con_datos.title = "Reporte"
        con_datos.append(["A", "B"])
        con_datos.append(["uno", 42])
        con_datos.add_table(Table(displayName="ConDatos", ref="A1:B2"))
        vacia = wb.create_sheet("REVISION IA")
        vacia.append(["C", "D"])
        vacia.add_table(Table(displayName="SinDatos", ref="A1:B1"))
        original = BytesIO()
        wb.save(original)

        reparado, retiradas = reparar_tablas_vacias(original.getvalue())
        self.assertEqual(len(retiradas), 1)
        with ZipFile(BytesIO(reparado)) as paquete:
            self.assertIsNone(paquete.testzip())
            self.assertEqual(
                [n for n in paquete.namelist() if n.startswith("xl/tables/")],
                ["xl/tables/table1.xml"],
            )
        antes = load_workbook(BytesIO(original.getvalue()))
        despues = load_workbook(BytesIO(reparado))
        self.assertEqual(list(antes["Reporte"].values), list(despues["Reporte"].values))
        self.assertEqual(list(antes["REVISION IA"].values), list(despues["REVISION IA"].values))
        self.assertEqual(len(despues["Reporte"].tables), 1)
        self.assertEqual(len(despues["REVISION IA"].tables), 0)
        self.assertEqual(reparar_tablas_vacias(reparado), (reparado, ()))


if __name__ == "__main__":
    unittest.main()
