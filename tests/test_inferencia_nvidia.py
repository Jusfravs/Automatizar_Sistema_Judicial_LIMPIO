"""Contratos de privacidad, evidencia e historial del auditor NVIDIA."""

import io
import json
import sqlite3
import tempfile
import unittest
from contextlib import closing
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd

from src.gestor_casos import GestorCasos
from src.gestor_cola import GestorCola
from src.historial_ia import normalizar_actuaciones, registrar_auditoria
from src.motor_temporal import reconstruir_cronologia
from src.proveedor_nvidia import (
    ErrorAuditoriaIA, auditar_nvidia, construir_contexto, validar_decision,
)
from scripts.auditar_inferencia_nvidia import ejecutar
from scripts.importar_revision_ia import _version_vigente, importar
from src.reporte_ia import preparar_vistas_ia


class AuditoriaNvidiaTests(unittest.TestCase):
    def setUp(self):
        self.datos = {
            "HISTORIAL_ACTUACIONES": [
                {"fecha": "01/02/2024", "tipo": "CITACION", "detalle":
                 "Citación a JUAN PEREZ, cédula 1712345678 en causa 07331-2024-00277"},
                {"fecha": "04/12/2025", "tipo": "OFICIO", "detalle":
                 "OFICIO: expediente de JUAN PEREZ, correo jp@example.com"},
            ],
            "ULTIMA ETAPA": "2 CITACION", "ULTIMA FASE": "2.1 CITACION",
            "ETAPA ACTUAL": "2 CITACION", "FASE ACTUAL": "2.1 CITACION",
            "eta_id ULTIMA ETAPA": 11, "fas_id ULTIMA FASE": 80,
            "eta_id ETAPA ACTUAL": 11, "fas_id FASE ACTUAL": 80,
            "FECHA FIN ULTIMA FASE": "04/12/2025",
        }
        self.actuaciones = normalizar_actuaciones(
            "07331-2024-00277", self.datos["HISTORIAL_ACTUACIONES"]
        )

    def test_contexto_no_contiene_identificadores_ni_texto_libre(self):
        contexto, mapa = construir_contexto(self.datos, self.actuaciones)
        payload = json.dumps(contexto, ensure_ascii=False)
        for prohibido in ("JUAN", "PEREZ", "1712345678", "07331-2024-00277",
                          "jp@example.com"):
            self.assertNotIn(prohibido, payload)
        self.assertEqual(len(mapa), 2)
        self.assertEqual(contexto["cronologia"][-1]["fecha"], "2025-12-04")

    def test_contexto_recupera_ids_de_resultados_historicos_sin_columnas_id(self):
        datos = {clave: valor for clave, valor in self.datos.items()
                 if not clave.startswith(("eta_id ", "fas_id "))}
        datos.update({
            "ULTIMA ETAPA": "6 LIQUIDACION Y EMBARGO",
            "ULTIMA FASE": "6.1 LIQUIDACION PERITO LIQUIDADOR",
            "ETAPA ACTUAL": "6 LIQUIDACION Y EMBARGO",
            "FASE ACTUAL": "6.2 MANDAMIENTO DE EJECUCION",
        })
        contexto, _ = construir_contexto(datos, self.actuaciones)
        self.assertEqual(contexto["inferencia_sistema"]["ultimo_hito"], {
            "eta_id": 15, "fas_id": 88, "fecha": "2025-12-04",
        })
        self.assertEqual(contexto["inferencia_sistema"]["estado_actual"], {
            "eta_id": 15, "fas_id": 89,
        })
        self.assertEqual(contexto["version"], "auditoria-nvidia-3")

    def test_validador_rechaza_correccion_sin_cambio_y_conservar_con_cambio(self):
        contexto, mapa = construir_contexto(self.datos, self.actuaciones)
        decision = {
            "decision": "CORREGIR",
            "estado_actual": {"eta_id": 11, "fas_id": 80},
            "ultimo_hito": {"eta_id": 11, "fas_id": 80},
            "evidencias": ["A1"], "confianza": 0.95,
            "motivo": "Solicita un cambio", "requiere_revision_humana": False,
        }
        with self.assertRaisesRegex(ErrorAuditoriaIA, "CORRECCION_SIN_CAMBIO"):
            validar_decision(decision, mapa, contexto)
        decision["decision"] = "CONSERVAR"
        decision["estado_actual"] = {"eta_id": 12, "fas_id": 82}
        with self.assertRaisesRegex(ErrorAuditoriaIA, "CONSERVAR_CON_CAMBIO"):
            validar_decision(decision, mapa, contexto)
        decision["estado_actual"] = {"eta_id": 11, "fas_id": 80}
        self.assertEqual(validar_decision(decision, mapa, contexto)["decision"],
                         "CONSERVAR")

    def test_mencion_posterior_de_fase_antigua_no_retrocede_cronologia(self):
        lectura = reconstruir_cronologia([
            {"ref": "A1", "fecha": "2024-01-01", "senales": ["CITACION"], "tipo": []},
            {"ref": "A2", "fecha": "2025-01-01", "senales": ["CONTESTACION"], "tipo": []},
            {"ref": "A3", "fecha": "2026-01-01", "senales": ["CITACION"], "tipo": []},
        ])
        self.assertEqual(lectura["ultimo_hito_candidato"]["ref"], "A2")
        self.assertEqual(lectura["ultimo_hito_candidato"]["fas_id"], 82)
        self.assertEqual(lectura["menciones_historicas"][0]["ref"], "A3")

    def test_validador_rechaza_evidencia_inventada_e_ids_incompatibles(self):
        _, mapa = construir_contexto(self.datos, self.actuaciones)
        decision = {
            "decision": "CORREGIR", "estado_actual": {"eta_id": 11, "fas_id": 80},
            "ultimo_hito": {"eta_id": 11, "fas_id": 80},
            "evidencias": ["A999"], "confianza": 0.98,
            "motivo": "Actuación posterior", "requiere_revision_humana": True,
        }
        with self.assertRaises(ErrorAuditoriaIA):
            validar_decision(decision, mapa)
        decision["evidencias"] = ["A2"]
        decision["estado_actual"] = {"eta_id": 12, "fas_id": 80}
        with self.assertRaises(ErrorAuditoriaIA):
            validar_decision(decision, mapa)
        decision["estado_actual"] = {"eta_id": 11, "fas_id": 80}
        self.assertEqual(validar_decision(decision, mapa)["evidencias"],
                         [self.actuaciones[1]["actuacion_id"]])

    def test_peticion_usa_clave_del_entorno_y_valida_json(self):
        contexto, mapa = construir_contexto(self.datos, self.actuaciones)
        decision = {
            "decision": "CONSERVAR", "estado_actual": {"eta_id": 11, "fas_id": 80},
            "ultimo_hito": {"eta_id": 11, "fas_id": 80},
            "evidencias": ["A1"], "confianza": 0.91,
            "motivo": "La citación está documentada", "requiere_revision_humana": False,
        }
        respuesta = {"choices": [{"message": {"content": json.dumps(decision)}}],
                     "usage": {"prompt_tokens": 123, "completion_tokens": 45}}
        capturado = {}

        def urlopen_falso(solicitud, timeout):
            capturado["solicitud"] = solicitud
            return io.BytesIO(json.dumps(respuesta).encode("utf-8"))

        with patch.dict("os.environ", {"NVIDIA_API_KEY": "clave-de-prueba"}), \
             patch("src.proveedor_nvidia.urlopen", side_effect=urlopen_falso):
            validada, consumo = auditar_nvidia(contexto, mapa)
        self.assertEqual(validada["decision"], "CONSERVAR")
        self.assertEqual(consumo["prompt_tokens"], 123)
        self.assertEqual(capturado["solicitud"].get_header("Authorization"),
                         "Bearer clave-de-prueba")
        self.assertNotIn(b"JUAN", capturado["solicitud"].data)


class HistorialPersistidoTests(unittest.TestCase):
    def test_piloto_con_contexto_obsoleto_no_se_ofrece_para_revision(self):
        with tempfile.TemporaryDirectory() as temporal:
            ruta = Path(temporal) / "casos.db"
            causa = "CAUSA-ANTIGUA"
            GestorCola(str(ruta))
            with closing(sqlite3.connect(ruta)) as conn:
                registrar_auditoria(
                    conn, numero_causa=causa, huella="anterior",
                    modelo="nvidia/test", modo="sombra", estado="REVISION",
                    decision={"decision": "CORREGIR", "confianza": 0.95,
                              "version_prompt": "auditoria-nvidia-2"},
                )
                conn.commit()
            reporte = pd.DataFrame({"NUMERO_JUICIO": [causa]})
            vista, hitos, revisiones = preparar_vistas_ia(ruta, reporte)
            self.assertEqual(vista.loc[0, "ESTADO AUDITORIA IA"], "PENDIENTE")
            self.assertTrue(hitos.empty)
            self.assertTrue(revisiones.empty)
            self.assertFalse(_version_vigente({"version_prompt": "auditoria-nvidia-1"}))
            self.assertFalse(_version_vigente({"version_prompt": "auditoria-nvidia-2"}))
            self.assertTrue(_version_vigente({"version_prompt": "auditoria-nvidia-3"}))

    def test_auditor_elige_postgres_cuando_configurado(self):
        with tempfile.TemporaryDirectory() as temporal:
            config = Path(temporal) / "config.json"
            config.write_text(json.dumps({
                "base_de_datos": {"motor": "postgres"},
                "inferencia_ia": {"modo": "sombra"},
            }), encoding="utf-8")
            with patch("src.auditoria_postgres.ejecutar_postgres",
                       return_value={"auditados": 0}) as ejecutar_pg:
                self.assertEqual(ejecutar(config, 1), {"auditados": 0})
            ejecutar_pg.assert_called_once()

    def test_registro_es_idempotente_y_no_cambia_fase_oficial(self):
        with tempfile.TemporaryDirectory() as temporal:
            ruta = Path(temporal) / "casos.db"
            gestor = GestorCola(str(ruta))
            causa = "07331-2024-00277"
            with closing(sqlite3.connect(ruta)) as conn:
                conn.execute("INSERT INTO juicios(numero_causa) VALUES (?)", (causa,))
                conn.commit()
            datos = {
                "ULTIMA ETAPA": "2 CITACION", "ULTIMA FASE": "2.1 CITACION",
                "ETAPA ACTUAL": "2 CITACION", "FASE ACTUAL": "2.1 CITACION",
                "HISTORIAL_ACTUACIONES": [
                    {"fecha": "04/12/2025", "tipo": "OFICIO", "detalle": "OFICIO"}
                ],
            }
            for _ in range(2):
                gestor.registrar_resultado_transaccional(
                    causa, {"datos": datos}, origen="PRUEBA"
                )
            with closing(sqlite3.connect(ruta)) as conn:
                self.assertEqual(conn.execute(
                    "SELECT COUNT(*) FROM actuaciones_procesales"
                ).fetchone()[0], 1)
                self.assertEqual(conn.execute(
                    "SELECT COUNT(*) FROM ejecuciones_inferencia"
                ).fetchone()[0], 1)
                guardado = json.loads(conn.execute(
                    "SELECT datos_json FROM resultados_expediente WHERE numero_causa=?",
                    (causa,),
                ).fetchone()[0])
            self.assertEqual(guardado["datos"]["FASE ACTUAL"], "2.1 CITACION")

    def test_lotes_reanudables_avanzan_a_la_siguiente_causa(self):
        with tempfile.TemporaryDirectory() as temporal:
            ruta = Path(temporal) / "casos.db"
            config = Path(temporal) / "config.json"
            config.write_text(json.dumps({
                "rutas": {"archivo_db": str(ruta)},
                "inferencia_ia": {"modo": "sombra"},
            }), encoding="utf-8")
            gestor = GestorCola(str(ruta))
            with closing(sqlite3.connect(ruta)) as conn:
                conn.executemany("INSERT INTO juicios(numero_causa) VALUES (?)",
                                 [("CAUSA-1",), ("CAUSA-2",)])
                conn.commit()
            for causa in ("CAUSA-1", "CAUSA-2"):
                gestor.registrar_resultado_transaccional(causa, {"datos": {
                    "ETAPA ACTUAL": "2 CITACION", "FASE ACTUAL": "2.1 CITACION",
                    "ULTIMA ETAPA": "2 CITACION", "ULTIMA FASE": "2.1 CITACION",
                    "HISTORIAL_ACTUACIONES": [{
                        "fecha": "04/12/2025", "detalle": "OFICIO",
                    }],
                }}, origen="PRUEBA")
            respuesta = ({
                "decision": "CONSERVAR",
                "estado_actual": {"eta_id": 11, "fas_id": 80},
                "evidencias": [], "confianza": 0.9,
                "motivo": "Coincide", "requiere_revision_humana": False,
            }, {"prompt_tokens": 10})
            with patch("scripts.auditar_inferencia_nvidia.auditar_nvidia",
                       return_value=respuesta) as llamada:
                primero = ejecutar(config, 1)
                segundo = ejecutar(config, 1)
            self.assertEqual(primero["auditados"], 1)
            self.assertEqual(segundo["auditados"], 1)
            self.assertEqual(llamada.call_count, 2)
            with closing(sqlite3.connect(ruta)) as conn:
                self.assertEqual(conn.execute(
                    "SELECT COUNT(*) FROM auditorias_ia"
                ).fetchone()[0], 2)

    def test_excel_agrega_revision_sin_contaminar_para_carga(self):
        with tempfile.TemporaryDirectory() as temporal:
            ruta_db = Path(temporal) / "casos.db"
            ruta_excel = Path(temporal) / "reporte.xlsx"
            cola = GestorCola(str(ruta_db))
            causa = "CAUSA-1"
            with closing(sqlite3.connect(ruta_db)) as conn:
                conn.execute("INSERT INTO juicios(numero_causa) VALUES (?)", (causa,))
                conn.commit()
            cola.registrar_resultado_transaccional(causa, {"datos": {
                "ULTIMA ETAPA": "2 CITACION", "ULTIMA FASE": "2.1 CITACION",
                "ETAPA ACTUAL": "2 CITACION", "FASE ACTUAL": "2.1 CITACION",
                "HISTORIAL_ACTUACIONES": [
                    {"fecha": "04/12/2025", "detalle": "OFICIO"}
                ],
            }}, origen="PRUEBA")
            with closing(sqlite3.connect(ruta_db)) as conn:
                actuacion_id = conn.execute(
                    "SELECT actuacion_id FROM actuaciones_procesales"
                ).fetchone()[0]
                registrar_auditoria(
                    conn, numero_causa=causa, huella="test", modelo="nvidia/test",
                    modo="sombra", estado="REVISION", decision={
                        "decision": "CORREGIR",
                        "estado_actual": {"eta_id": 12, "fas_id": 82},
                        "evidencias": [actuacion_id], "confianza": 0.98,
                        "motivo": "Actuación posterior",
                        "requiere_revision_humana": True,
                    },
                )
                conn.execute("""
                    INSERT INTO hitos_procesales
                    (numero_causa, actuacion_id, eta_id, fas_id, fuente,
                     condicion, version) VALUES (?, ?, 12, 82, 'NVIDIA',
                     'PROPUESTA', 'test')
                """, (causa, actuacion_id))
                conn.commit()
            gestor = GestorCasos.__new__(GestorCasos)
            gestor.ruta_final = str(ruta_excel)
            gestor.ruta_config = str(Path(temporal) / "config.json")
            gestor.config = {
                "rutas": {"archivo_db": str(ruta_db)},
                "inferencia_ia": {"modo": "sombra"},
            }
            Path(gestor.ruta_config).write_text(
                json.dumps(gestor.config), encoding="utf-8"
            )
            gestor.df = pd.DataFrame({
                "CODIGO_JUICIO": [1], "NUMERO_JUICIO": [causa],
                "ULTIMA ETAPA": ["2 CITACION"], "ULTIMA FASE": ["2.1 CITACION"],
                "ETAPA ACTUAL": ["2 CITACION"], "FASE ACTUAL": ["2.1 CITACION"],
                "FECHA FIN ULTIMA FASE": ["04/12/2025"],
                "FECHA INICIO FASE ACTUAL": ["04/12/2025"],
                "COMENTARIO_ULTIMO": [""],
            })
            gestor.exportar_excel(datetime(2026, 9, 24))
            libro = openpyxl.load_workbook(ruta_excel, read_only=True)
            self.assertEqual(set(libro.sheetnames), {
                "Reporte", "PARA CARGA", "HISTORIAL HITOS", "REVISION IA"
            })
            carga = pd.read_excel(ruta_excel, sheet_name="PARA CARGA")
            reporte = pd.read_excel(
                ruta_excel, sheet_name="Reporte",
                header=GestorCasos.FILA_ENCABEZADO_REPORTE - 1,
            )
            self.assertNotIn("ESTADO AUDITORIA IA", carga.columns)
            self.assertEqual(reporte.loc[0, "ESTADO AUDITORIA IA"], "REVISION")
            self.assertEqual(libro["REVISION IA"].max_row, 2)
            libro.close()
            editable = openpyxl.load_workbook(ruta_excel)
            hoja = editable["REVISION IA"]
            encabezados = {
                celda.value: celda.column for celda in hoja[1] if celda.value
            }
            hoja.cell(2, encabezados["DECISION HUMANA"]).value = "ACEPTAR_SISTEMA"
            editable.save(ruta_excel)
            editable.close()
            self.assertEqual(importar(ruta_excel, Path(gestor.ruta_config)), 1)
            with closing(sqlite3.connect(ruta_db)) as conn:
                self.assertEqual(conn.execute(
                    "SELECT estado, decision_humana FROM revisiones_ia"
                ).fetchone(), ("RESUELTA", "ACEPTAR_SISTEMA"))
                oficial = json.loads(conn.execute(
                    "SELECT datos_json FROM resultados_expediente"
                ).fetchone()[0])
            self.assertEqual(oficial["datos"]["FASE ACTUAL"], "2.1 CITACION")
            with self.assertRaises(ValueError):
                importar(ruta_excel, Path(gestor.ruta_config))


if __name__ == "__main__":
    unittest.main()
