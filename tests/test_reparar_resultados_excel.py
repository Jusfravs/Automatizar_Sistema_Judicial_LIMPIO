import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from scripts import reparar_resultados_excel as reparacion


class StorageFalso:
    def __init__(self, objetos):
        self.objetos = dict(objetos)
        self.subidas = []

    def descargar(self, ruta):
        if ruta not in self.objetos:
            raise RuntimeError("STORAGE_DESCARGA_HTTP_404")
        return self.objetos[ruta]

    def subir(self, ruta, contenido):
        self.objetos[ruta] = contenido
        self.subidas.append(ruta)


def libro(*, tabla_vacia):
    wb = Workbook()
    hoja = wb.active
    hoja.append(["A", "B"])
    if tabla_vacia:
        hoja.add_table(Table(displayName="SinDatos", ref="A1:B1"))
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


class RepararResultadosTests(unittest.TestCase):
    def test_enumera_completos_incluso_sin_resultado_principal(self):
        cursor = Mock()
        cursor.fetchone.return_value = (0,)
        cursor.fetchall.return_value = [
            ("uno", "lote", "b2:entradas/uno.xlsx", "b2:resultados/uno.xlsx", ""),
            ("dos", "solo", "b2:entradas/dos.xlsx", None,
             "EXCEL_LOTE_NO_DISPONIBLE"),
            ("tres", "pendientes", "b2:entradas/tres.xlsx",
             "b2:resultados/tres.xlsx", ""),
            ("cuatro", "lote", "entradas/cuatro.xlsx",
             "resultados/cuatro.xlsx", ""),
        ]
        conexion = Mock()
        conexion.cursor.return_value = cursor
        repo = Mock()
        repo._get_connection.return_value = conexion
        with patch.object(reparacion, "cargar_comun", return_value={}), \
                patch.object(reparacion, "_repositorio", return_value=repo):
            rutas = reparacion._rutas(Path("config.json"), comprobar_sin_lotes=True)

        self.assertEqual(rutas, [
            ("b2:resultados/uno.xlsx", False),
            ("b2:resultados/uno_completo.xlsx", True),
            ("b2:resultados/dos_completo.xlsx", True),
            ("b2:resultados/tres.xlsx", False),
            ("resultados/cuatro.xlsx", False),
            ("resultados/cuatro_completo.xlsx", True),
        ])
        conexion.rollback.assert_called_once()
        conexion.close.assert_called_once()

    def test_repara_completo_existente_y_omite_solo_el_404_opcional(self):
        roto = libro(tabla_vacia=True)
        sano = libro(tabla_vacia=False)
        completo = "b2:resultados/uno_completo.xlsx"
        storage = StorageFalso({
            "b2:resultados/uno.xlsx": sano,
            completo: roto,
            "resultados/dos.xlsx": sano,
        })
        rutas = [
            ("b2:resultados/uno.xlsx", False),
            (completo, True),
            ("resultados/dos.xlsx", False),
            ("resultados/dos_completo.xlsx", True),
        ]
        with tempfile.TemporaryDirectory() as temporal:
            respaldo = Path(temporal) / "respaldo"
            with patch.object(reparacion, "_rutas", return_value=rutas) as listar, \
                    patch.object(reparacion.ClienteStorage, "desde_entorno",
                                 return_value=storage):
                resumen = reparacion.ejecutar(Path("config.json"), True, respaldo, 3)
            listar.assert_called_once_with(Path("config.json"), comprobar_sin_lotes=True)
            manifiesto = json.loads((respaldo / "manifiesto.json").read_text(encoding="utf-8"))
            self.assertEqual(manifiesto[0]["ruta"], completo)
            self.assertEqual((respaldo / manifiesto[0]["respaldo"]).read_bytes(), roto)

        self.assertEqual(resumen["resultados"], 3)
        self.assertEqual(resumen["completos"], 1)
        self.assertEqual(resumen["completos_ausentes"], 1)
        self.assertEqual(resumen["afectados"], 1)
        self.assertTrue(resumen["aplicado"])
        self.assertEqual(storage.subidas, [completo])
        self.assertEqual(len(load_workbook(BytesIO(storage.objetos[completo])).active.tables), 0)

    def test_un_error_distinto_de_404_no_se_considera_ausencia(self):
        error = RuntimeError("STORAGE_DESCARGA_HTTP_403")
        self.assertFalse(reparacion._objeto_ausente(error))
        error_b2 = RuntimeError("NoSuchKey")
        error_b2.response = {
            "ResponseMetadata": {"HTTPStatusCode": 404},
            "Error": {"Code": "NoSuchKey"},
        }
        self.assertTrue(reparacion._objeto_ausente(error_b2))


if __name__ == "__main__":
    unittest.main()
