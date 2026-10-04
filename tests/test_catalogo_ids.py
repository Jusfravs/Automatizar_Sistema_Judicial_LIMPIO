import os
import tempfile
import unittest
from datetime import datetime

import pandas as pd
import openpyxl

from src.agente_extractor import ResultadoInferencia
from src.catalogo_procesal import (
    ETAPAS,
    FASES,
    canonicalizar_fase,
    id_etapa,
    id_fase,
    validar_ids,
)
from src.gestor_casos import GestorCasos


class CatalogoProcesalTests(unittest.TestCase):
    def test_catalogo_oficial_y_aliases(self):
        self.assertEqual(list(ETAPAS.values()), [10, 11, 12, 13, 14, 15])
        self.assertEqual(list(FASES.values()), list(range(77, 93)))
        self.assertEqual(id_etapa("1 PRESENTACION Y CALIFICACION"), 10)
        self.assertEqual(id_etapa("6 LIQUIDACION Y EMBARGO"), 15)
        self.assertEqual(id_fase("1.1 PRESENTAR DEMANDA"), 77)
        self.assertEqual(id_fase("6.5 CONGELAMIENTO DE CUENTAS / CIERRE"), 92)
        self.assertEqual(id_fase("CONTESTACION"), 82)
        self.assertEqual(
            canonicalizar_fase("4.2 AUDIENCIA / ACTA RESUMEN"),
            "4.2 AUDIENCIA",
        )
        self.assertIsNone(id_fase("4.3 ACUERDO DE MEDIACION"))

    def test_resultado_inferencia_incluye_los_cuatro_ids(self):
        resultado = ResultadoInferencia(
            "1 PRESENTACION Y CALIFICACION",
            "1.3 CALIFICACION",
            "01/09/2026",
            etapa_actual="2 CITACION",
            fase_actual="2.1 CITACION (PERSONA/BOLETA)",
        )
        datos = {
            "ULTIMA ETAPA": resultado.ultima_etapa,
            "ULTIMA FASE": resultado.ultima_fase,
            "ETAPA ACTUAL": resultado.etapa_actual,
            "FASE ACTUAL": resultado.fase_actual,
            "eta_id ULTIMA ETAPA": resultado.eta_id_ultima_etapa,
            "fas_id ULTIMA FASE": resultado.fas_id_ultima_fase,
            "eta_id ETAPA ACTUAL": resultado.eta_id_etapa_actual,
            "fas_id FASE ACTUAL": resultado.fas_id_fase_actual,
        }
        self.assertTrue(validar_ids(datos))
        self.assertEqual(resultado.fas_id_fase_actual, 80)


class ExportacionSistemasTests(unittest.TestCase):
    def _gestor(self):
        gestor = GestorCasos.__new__(GestorCasos)
        gestor.hoja = "origen"
        gestor.df = pd.DataFrame({
            "CODIGO_JUICIO": [1, 2, 3, 4],
            "NUMERO_JUICIO": ["CAUSA-1", "CAUSA-1", "CAUSA-3", "CAUSA-4"],
            "COMENTARIO_ULTIMO": [
                "", "", "REVISION MANUAL: La consulta no devolvio resultados", "REVISION MANUAL",
            ],
            "FECHA INICIO JUICIO": ["01/01/2026"] * 4,
            "FECHA FIN ULTIMA FASE": ["01/09/2026"] * 4,
            "ULTIMA ETAPA": [
                "1 PRESENTACION Y CALIFICACION", "2 CITACION",
                None, "4 AUDIENCIA",
            ],
            "ULTIMA FASE": [
                "1.3 CALIFICACION", "2.1 CITACION (PERSONA/BOLETA)",
                None, "4.3 ACUERDO DE MEDIACION",
            ],
            "FECHA INICIO FASE ACTUAL": ["16/09/2026"] * 4,
            "ETAPA ACTUAL": [
                "2 CITACION", "CONTESTACION", None, "REVISION MANUAL",
            ],
            "FASE ACTUAL": [
                "2.1 CITACION (PERSONA/BOLETA)", "CONTESTACION",
                None, "REVISION MANUAL",
            ],
        })
        return gestor

    def test_exporta_reporte_y_para_carga_con_dias_consistentes(self):
        gestor = self._gestor()
        with tempfile.TemporaryDirectory() as temporal:
            gestor.ruta_final = os.path.join(temporal, "reporte.xlsx")
            gestor.exportar_excel(datetime(2026, 9, 18))
            hojas = {
                "Reporte": pd.read_excel(
                    gestor.ruta_final,
                    sheet_name="Reporte",
                    header=GestorCasos.FILA_ENCABEZADO_REPORTE - 1,
                ),
                "PARA CARGA": pd.read_excel(
                    gestor.ruta_final, sheet_name="PARA CARGA"
                ),
            }
            wb = openpyxl.load_workbook(gestor.ruta_final)
            ws_reporte = wb["Reporte"]
            ws_carga = wb["PARA CARGA"]
            self.assertEqual(ws_reporte["A1"].value, "REPORTE FINAL DE GESTIÓN JUDICIAL")
            self.assertEqual(ws_reporte["A3"].value, "Total de registros")
            self.assertEqual(ws_reporte["A4"].value, 4)
            indicadores = {
                celda.value: ws_reporte.cell(4, celda.column).value
                for celda in ws_reporte[3] if celda.value
            }
            self.assertEqual(indicadores["Listos para carga"], 3)
            self.assertEqual(indicadores["Observados / no cargables"], 1)
            self.assertEqual(indicadores["Revisión manual / ID pendiente"], 1)
            self.assertEqual(ws_reporte.freeze_panes, "A7")
            self.assertEqual(ws_carga.freeze_panes, "A2")
            self.assertFalse(ws_reporte.sheet_view.showGridLines)
            self.assertFalse(ws_carga.sheet_view.showGridLines)
            self.assertIn("TablaReporteJudicial", ws_reporte.tables)
            self.assertIn("TablaParaCarga", ws_carga.tables)
            self.assertIsNone(ws_reporte.auto_filter.ref)
            self.assertIsNone(ws_carga.auto_filter.ref)
            self.assertEqual(ws_reporte["A6"].value, "CODIGO_JUICIO")
            self.assertEqual(ws_carga["A1"].value, "CODIGO_JUICIO")
            self.assertEqual(ws_reporte["A9"].fill.fgColor.rgb, "FFFCE8E6")
            wb.close()

        self.assertEqual(list(hojas), ["Reporte", "PARA CARGA"])
        self.assertEqual(len(hojas["Reporte"]), 4)
        self.assertEqual(len(hojas["PARA CARGA"]), 3)
        self.assertNotIn(" ", hojas["Reporte"].columns)
        self.assertEqual(hojas["Reporte"].loc[0, "DIAS TRANSCURRIDOS"], 2)
        self.assertEqual(hojas["PARA CARGA"].loc[0, "DIAS TRANSCURRIDOS"], 2)
        self.assertEqual(hojas["Reporte"].loc[0, "fas_id FASE ACTUAL"], 80)
        self.assertEqual(hojas["Reporte"].loc[1, "fas_id FASE ACTUAL"], 82)
        self.assertEqual(hojas["Reporte"].loc[2, "FASE ACTUAL"], "MAL INGRESADO")
        self.assertTrue(pd.isna(hojas["Reporte"].loc[3, "fas_id ULTIMA FASE"]))
        self.assertIn(
            "ID_NO_CATALOGADO",
            hojas["Reporte"].loc[3, "COMENTARIO_ULTIMO"],
        )

    def test_rechaza_codigo_juicio_duplicado(self):
        gestor = self._gestor()
        gestor.df.loc[1, "CODIGO_JUICIO"] = 1
        with tempfile.TemporaryDirectory() as temporal:
            gestor.ruta_final = os.path.join(temporal, "reporte.xlsx")
            with self.assertRaisesRegex(ValueError, "CODIGO_JUICIO_DUPLICADO"):
                gestor.exportar_excel(datetime(2026, 9, 18))

    def test_excluye_estado_no_correspondiente_a_cartera(self):
        gestor = self._gestor()
        gestor.df.loc[3, "ETAPA ACTUAL"] = "EXCLUIDO_NO_CORRESPONDE"
        gestor.df.loc[3, "FASE ACTUAL"] = "EXCLUIDO_NO_CORRESPONDE"
        reporte, para_carga = gestor._preparar_exportacion(
            datetime(2026, 9, 18)
        )
        self.assertEqual(len(reporte), 4)
        self.assertEqual(len(para_carga), 2)
        self.assertEqual(reporte.loc[3, "ETAPA ACTUAL"], "MAL INGRESADO")


if __name__ == "__main__":
    unittest.main()
