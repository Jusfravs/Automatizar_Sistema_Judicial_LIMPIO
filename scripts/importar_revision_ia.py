"""Importa decisiones de la hoja REVISION IA sin modificar la clasificación oficial."""

from __future__ import annotations

import argparse
import json
import sqlite3
from contextlib import closing
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from src.catalogo_procesal import ETAPAS, FASES
from src.proveedor_nvidia import VERSION_PROMPT


DECISIONES = {"ACEPTAR_SISTEMA", "ACEPTAR_IA", "CORREGIR_MANUALMENTE", "POSPONER"}


def _entero(valor):
    if isinstance(valor, bool) or valor in (None, ""):
        return None
    try:
        numero = int(valor)
    except (TypeError, ValueError):
        return None
    return numero if str(valor).strip() in {str(numero), f"{numero}.0"} else None


def _par_valido(eta_id, fas_id):
    if eta_id not in ETAPAS.values() or fas_id not in FASES.values():
        return False
    fase = next(nombre for nombre, clave in FASES.items() if clave == fas_id)
    return fase.startswith(str(eta_id - 9) + ".")


def _version_vigente(propuesta):
    version = str(propuesta.get("version_prompt") or "")
    return not version.startswith("auditoria-nvidia-") or version == VERSION_PROMPT


def _filas_revision(ruta_excel: Path):
    libro = load_workbook(ruta_excel, read_only=True, data_only=True)
    try:
        if "REVISION IA" not in libro:
            raise ValueError("HOJA_REVISION_IA_AUSENTE")
        hoja = libro["REVISION IA"]
        filas = hoja.iter_rows(values_only=True)
        encabezados = [str(celda or "").strip() for celda in next(filas)]
        requeridos = {"AUDITORIA_ID", "NUMERO_JUICIO", "DECISION HUMANA",
                      "eta_id MANUAL", "fas_id MANUAL", "OBSERVACION HUMANA"}
        if not requeridos <= set(encabezados):
            raise ValueError("COLUMNAS_REVISION_IA_INCOMPLETAS")
        for numero_fila, valores in enumerate(filas, start=2):
            fila = dict(zip(encabezados, valores))
            if str(fila.get("DECISION HUMANA") or "").strip():
                yield numero_fila, fila
    finally:
        libro.close()


def importar(ruta_excel: Path, ruta_config: Path, simular: bool = False) -> int:
    config = json.loads(ruta_config.read_text(encoding="utf-8-sig"))
    decisiones = list(_filas_revision(ruta_excel))
    if config.get("base_de_datos", {}).get("motor") == "postgres":
        return _importar_postgres(config, decisiones, simular)
    ruta_db = Path(config["rutas"]["archivo_db"])
    if not ruta_db.is_absolute():
        ruta_db = ruta_config.parent / ruta_db
    if not ruta_db.is_file():
        raise FileNotFoundError(ruta_db)
    vistos = set()
    actualizaciones = []
    with closing(sqlite3.connect(ruta_db, timeout=30)) as conn:
        if not simular:
            conn.execute("BEGIN IMMEDIATE")
        for numero_fila, fila in decisiones:
            auditoria_id = _entero(fila["AUDITORIA_ID"])
            causa = str(fila["NUMERO_JUICIO"] or "").strip()
            decision = str(fila["DECISION HUMANA"] or "").strip().upper()
            eta_id = _entero(fila["eta_id MANUAL"])
            fas_id = _entero(fila["fas_id MANUAL"])
            observacion = str(fila["OBSERVACION HUMANA"] or "").strip()
            if not auditoria_id or not causa or decision not in DECISIONES:
                raise ValueError(f"FILA_{numero_fila}_DECISION_INVALIDA")
            if auditoria_id in vistos:
                raise ValueError(f"FILA_{numero_fila}_AUDITORIA_DUPLICADA")
            vistos.add(auditoria_id)
            registro = conn.execute("""
                SELECT a.numero_causa, a.estado, a.decision_json, r.estado
                FROM auditorias_ia a JOIN revisiones_ia r ON r.auditoria_id=a.id
                WHERE a.id=?
            """, (auditoria_id,)).fetchone()
            ultimo = conn.execute("""
                SELECT MAX(id) FROM auditorias_ia WHERE numero_causa=?
            """, (causa,)).fetchone()[0]
            if (not registro or registro[0] != causa or ultimo != auditoria_id
                    or registro[1] not in {"REVISION", "INSUFICIENTE"}
                    or registro[3] != "PENDIENTE"):
                raise ValueError(f"FILA_{numero_fila}_REVISION_OBSOLETA")
            propuesta = json.loads(registro[2] or "{}")
            if not _version_vigente(propuesta):
                raise ValueError(f"FILA_{numero_fila}_REVISION_OBSOLETA")
            if decision == "ACEPTAR_IA":
                actual = propuesta.get("estado_actual") or {}
                if propuesta.get("decision") != "CORREGIR" or not _par_valido(
                    actual.get("eta_id"), actual.get("fas_id")
                ):
                    raise ValueError(f"FILA_{numero_fila}_PROPUESTA_IA_INVALIDA")
                evidencias = propuesta.get("evidencias") or []
                if not evidencias or any(not conn.execute("""
                    SELECT 1 FROM actuaciones_procesales
                    WHERE actuacion_id=? AND numero_causa=?
                """, (ref, causa)).fetchone() for ref in evidencias):
                    raise ValueError(f"FILA_{numero_fila}_EVIDENCIA_INVALIDA")
            if decision == "CORREGIR_MANUALMENTE" and (
                not _par_valido(eta_id, fas_id) or not observacion
            ):
                raise ValueError(f"FILA_{numero_fila}_CORRECCION_MANUAL_INVALIDA")
            if decision != "CORREGIR_MANUALMENTE" and (eta_id or fas_id):
                raise ValueError(f"FILA_{numero_fila}_IDS_MANUALES_INESPERADOS")
            actualizaciones.append((
                "PENDIENTE" if decision == "POSPONER" else "RESUELTA",
                decision, eta_id, fas_id, observacion,
                datetime.now(timezone.utc).isoformat() if decision != "POSPONER" else None,
                auditoria_id,
            ))
        if not simular:
            with conn:
                conn.executemany("""
                    UPDATE revisiones_ia SET estado=?, decision_humana=?,
                    eta_id_manual=?, fas_id_manual=?, observacion=?, revisado_en=?
                    WHERE auditoria_id=? AND estado='PENDIENTE'
                """, actualizaciones)
    return len(actualizaciones)


def _importar_postgres(config: dict, decisiones: list, simular: bool) -> int:
    from src.repositorio_postgres import RepositorioColaPostgres

    repo = RepositorioColaPostgres.desde_config(config.get("base_de_datos"))
    conn = repo._get_connection()
    vistos = set()
    actualizaciones = []
    try:
        with conn.cursor() as cursor:
            for numero_fila, fila in decisiones:
                auditoria_id = _entero(fila["AUDITORIA_ID"])
                causa = str(fila["NUMERO_JUICIO"] or "").strip()
                decision = str(fila["DECISION HUMANA"] or "").strip().upper()
                eta_id = _entero(fila["eta_id MANUAL"])
                fas_id = _entero(fila["fas_id MANUAL"])
                observacion = str(fila["OBSERVACION HUMANA"] or "").strip()
                if not auditoria_id or not causa or decision not in DECISIONES:
                    raise ValueError(f"FILA_{numero_fila}_DECISION_INVALIDA")
                if auditoria_id in vistos:
                    raise ValueError(f"FILA_{numero_fila}_AUDITORIA_DUPLICADA")
                vistos.add(auditoria_id)
                cursor.execute("""
                    SELECT a.numero_causa, a.estado, a.decision_json, r.estado
                    FROM auditorias_ia a JOIN revisiones_ia r ON r.auditoria_id=a.id
                    WHERE a.id=%s
                """, (auditoria_id,))
                registro = cursor.fetchone()
                cursor.execute("""
                    SELECT MAX(id) FROM auditorias_ia WHERE numero_causa=%s
                """, (causa,))
                ultimo = cursor.fetchone()[0]
                if (not registro or registro[0] != causa or ultimo != auditoria_id
                        or registro[1] not in {"REVISION", "INSUFICIENTE"}
                        or registro[3] != "PENDIENTE"):
                    raise ValueError(f"FILA_{numero_fila}_REVISION_OBSOLETA")
                propuesta = registro[2] if isinstance(registro[2], dict) else json.loads(registro[2] or "{}")
                if not _version_vigente(propuesta):
                    raise ValueError(f"FILA_{numero_fila}_REVISION_OBSOLETA")
                if decision == "ACEPTAR_IA":
                    actual = propuesta.get("estado_actual") or {}
                    if propuesta.get("decision") != "CORREGIR" or not _par_valido(
                        actual.get("eta_id"), actual.get("fas_id")
                    ):
                        raise ValueError(f"FILA_{numero_fila}_PROPUESTA_IA_INVALIDA")
                    evidencias = propuesta.get("evidencias") or []
                    if not evidencias:
                        raise ValueError(f"FILA_{numero_fila}_EVIDENCIA_INVALIDA")
                    for ref in evidencias:
                        cursor.execute("""
                            SELECT 1 FROM actuaciones_procesales
                            WHERE actuacion_id=%s AND numero_causa=%s
                        """, (ref, causa))
                        if not cursor.fetchone():
                            raise ValueError(f"FILA_{numero_fila}_EVIDENCIA_INVALIDA")
                if decision == "CORREGIR_MANUALMENTE" and (
                    not _par_valido(eta_id, fas_id) or not observacion
                ):
                    raise ValueError(f"FILA_{numero_fila}_CORRECCION_MANUAL_INVALIDA")
                if decision != "CORREGIR_MANUALMENTE" and (eta_id or fas_id):
                    raise ValueError(f"FILA_{numero_fila}_IDS_MANUALES_INESPERADOS")
                actualizaciones.append((
                    "PENDIENTE" if decision == "POSPONER" else "RESUELTA",
                    decision, eta_id, fas_id, observacion,
                    datetime.now(timezone.utc) if decision != "POSPONER" else None,
                    auditoria_id,
                ))
            if not simular:
                cursor.executemany("""
                    UPDATE revisiones_ia SET estado=%s, decision_humana=%s,
                    eta_id_manual=%s, fas_id_manual=%s, observacion=%s,
                    revisado_en=%s
                    WHERE auditoria_id=%s AND estado='PENDIENTE'
                """, actualizaciones)
                conn.commit()
            else:
                conn.rollback()
        return len(actualizaciones)
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", type=Path, required=True)
    parser.add_argument("--config", type=Path, default=Path("config.json"))
    parser.add_argument("--simular", action="store_true")
    args = parser.parse_args()
    cantidad = importar(args.excel.resolve(), args.config.resolve(), args.simular)
    print(json.dumps({"revisiones": cantidad, "simulado": args.simular}))


if __name__ == "__main__":
    main()
