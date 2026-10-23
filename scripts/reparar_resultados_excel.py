"""Audita y repara resultados XLSX ya publicados en B2 y Supabase Storage.

Incluye el resultado principal y el reporte completo separado si existe.
Por defecto solo lee. Para reemplazar objetos exige --aplicar y --respaldo.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from consola import _repositorio  # noqa: E402
from src.configuracion_lotes import cargar_comun  # noqa: E402
from src.reparacion_excel import reparar_tablas_vacias  # noqa: E402
from src.servicio_portal import ClienteStorage  # noqa: E402


def _sha(contenido: bytes) -> str:
    return hashlib.sha256(contenido).hexdigest()


def _rutas(config: Path, *, comprobar_sin_lotes: bool) -> list[tuple[str, bool]]:
    """Lista principales y posibles completos; el segundo valor indica si es opcional."""
    repo = _repositorio(cargar_comun(config))
    conexion = repo._get_connection()
    try:
        cursor = conexion.cursor()
        cursor.execute("SET TRANSACTION READ ONLY")
        if comprobar_sin_lotes:
            cursor.execute(
                """SELECT count(*) FROM public.solicitudes_lote
                   WHERE estado IN ('SOLICITADA', 'TOMADA', 'PREPARANDO', 'EN_CURSO')"""
            )
            if cursor.fetchone()[0]:
                raise RuntimeError("HAY_LOTES_ACTIVOS")
        cursor.execute(
            """SELECT id, modo, archivo_ruta, resultado_ruta, mensaje
               FROM public.solicitudes_lote
               WHERE (resultado_ruta IS NOT NULL AND resultado_ruta <> '')
                  OR mensaje LIKE '%EXCEL_LOTE_NO_DISPONIBLE%'
               ORDER BY id"""
        )
        rutas = []
        for id_lote, modo, archivo_ruta, resultado_ruta, mensaje in cursor.fetchall():
            if resultado_ruta:
                rutas.append((resultado_ruta, False))
            if modo != "pendientes" and (
                resultado_ruta or "EXCEL_LOTE_NO_DISPONIBLE" in (mensaje or "")
            ):
                prefijo = "b2:" if archivo_ruta.startswith("b2:") else ""
                rutas.append((f"{prefijo}resultados/{id_lote}_completo.xlsx", True))
        return list(dict.fromkeys(rutas))
    finally:
        conexion.rollback()
        conexion.close()


def _verificar_celdas(original: bytes, reparado: bytes) -> None:
    antes = load_workbook(BytesIO(original), read_only=True, data_only=False)
    despues = load_workbook(BytesIO(reparado), read_only=True, data_only=False)
    try:
        if antes.sheetnames != despues.sheetnames:
            raise ValueError("XLSX_HOJAS_CAMBIARON")
        for nombre in antes.sheetnames:
            hoja_antes = antes[nombre]
            hoja_despues = despues[nombre]
            if (hoja_antes.max_row, hoja_antes.max_column) != (
                hoja_despues.max_row, hoja_despues.max_column
            ):
                raise ValueError("XLSX_DIMENSIONES_CAMBIARON")
            if any(a != b for a, b in zip(
                hoja_antes.values, hoja_despues.values, strict=True
            )):
                raise ValueError("XLSX_CELDAS_CAMBIARON")
    finally:
        antes.close()
        despues.close()


def _objeto_ausente(exc: Exception) -> bool:
    """Reconoce solo un 404 real; los errores de red y permisos abortan la auditoría."""
    if str(exc) == "STORAGE_DESCARGA_HTTP_404":
        return True
    respuesta = getattr(exc, "response", None)
    if not isinstance(respuesta, dict):
        return False
    estado = respuesta.get("ResponseMetadata", {}).get("HTTPStatusCode")
    codigo = respuesta.get("Error", {}).get("Code")
    return estado == 404 and codigo in {"404", "NoSuchKey", "NotFound"}


def ejecutar(config: Path, aplicar: bool, respaldo: Path | None, esperado: int | None) -> dict:
    rutas = _rutas(config, comprobar_sin_lotes=aplicar)
    storage = ClienteStorage.desde_entorno()
    preparados = []
    encontrados = []
    completos_ausentes = 0
    for ruta, opcional in rutas:
        if not ruta.endswith(".xlsx") or not (
            ruta.startswith("b2:resultados/") or ruta.startswith("resultados/")
        ):
            raise ValueError("RUTA_RESULTADO_INESPERADA")
        try:
            original = storage.descargar(ruta)
        except Exception as exc:
            if opcional and _objeto_ausente(exc):
                completos_ausentes += 1
                continue
            raise
        encontrados.append(ruta)
        reparado, eliminadas = reparar_tablas_vacias(original)
        if eliminadas:
            _verificar_celdas(original, reparado)
            preparados.append((ruta, original, reparado, eliminadas))

    if esperado is not None and len(encontrados) != esperado:
        raise RuntimeError(f"CANTIDAD_RESULTADOS_CAMBIO:{len(encontrados)}")

    resumen = {
        "resultados": len(encontrados),
        "principales": sum(not ruta.endswith("_completo.xlsx") for ruta in encontrados),
        "completos": sum(ruta.endswith("_completo.xlsx") for ruta in encontrados),
        "completos_ausentes": completos_ausentes,
        "afectados": len(preparados),
        "sin_cambios": len(encontrados) - len(preparados),
        "b2_afectados": sum(ruta.startswith("b2:") for ruta, *_ in preparados),
        "supabase_afectados": sum(not ruta.startswith("b2:") for ruta, *_ in preparados),
        "aplicado": False,
    }
    if not aplicar or not preparados:
        return resumen
    if respaldo is None:
        raise ValueError("RESPALDO_OBLIGATORIO")
    respaldo.mkdir(parents=True, exist_ok=False)
    manifiesto = []
    for ruta, original, reparado, eliminadas in preparados:
        archivo = hashlib.sha256(ruta.encode("utf-8")).hexdigest() + ".xlsx"
        (respaldo / archivo).write_bytes(original)
        manifiesto.append({
            "ruta": ruta,
            "respaldo": archivo,
            "sha256_original": _sha(original),
            "sha256_reparado": _sha(reparado),
            "tablas_retiradas": list(eliminadas),
        })
    (respaldo / "manifiesto.json").write_text(
        json.dumps(manifiesto, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    # Todos los originales ya estan en disco antes de reemplazar el primero.
    for ruta, original, reparado, _ in preparados:
        try:
            if _sha(storage.descargar(ruta)) != _sha(original):
                raise RuntimeError("RESULTADO_CAMBIO_DURANTE_LA_REPARACION")
            storage.subir(ruta, reparado)
            if _sha(storage.descargar(ruta)) != _sha(reparado):
                raise RuntimeError("RESULTADO_NO_VERIFICADO")
        except Exception:
            # Restaurar solo si el contenido remoto ya pudo haber cambiado.
            if _sha(storage.descargar(ruta)) != _sha(original):
                storage.subir(ruta, original)
                if _sha(storage.descargar(ruta)) != _sha(original):
                    raise RuntimeError("RESTAURACION_FALLIDA") from None
            raise
    resumen["aplicado"] = True
    resumen["respaldo"] = str(respaldo)
    return resumen


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=RAIZ / "config_supabase.json")
    parser.add_argument("--aplicar", action="store_true")
    parser.add_argument("--respaldo", type=Path)
    parser.add_argument("--esperados", type=int, help=(
        "Número de XLSX existentes: principales y completos; excluye completos ausentes"
    ))
    args = parser.parse_args()
    if args.aplicar and args.respaldo is None:
        parser.error("--aplicar requiere --respaldo")
    try:
        print(json.dumps(
            ejecutar(args.config, args.aplicar, args.respaldo, args.esperados),
            ensure_ascii=False,
        ))
        return 0
    except Exception as exc:
        # Las excepciones HTTP pueden contener cabeceras; solo se imprime el codigo.
        print(json.dumps({"ok": False, "codigo": str(exc).split(":")[0]}))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
