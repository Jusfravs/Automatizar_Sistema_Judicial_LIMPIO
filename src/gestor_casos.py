# src/gestor_casos.py
import os
import sys
import json
import shutil
import errno
import tempfile
import time
from datetime import datetime
import pandas as pd
from pandas.errors import EmptyDataError
from src.catalogo_procesal import (
    CAMPOS_ID,
    enriquecer_datos_procesales,
)
from src.logger_config import obtener_logger
from src.ultima_gestion import (
    MARCA_REVISION_ULTIMA_GESTION,
    MOTIVO_REVISION_ULTIMA_GESTION,
    CAMPO_ESTADO_ULTIMA_GESTION,
    CAMPO_FECHA_ULTIMA_GESTION,
    enriquecer_ultima_gestion_judicial,
)
from src.reporte_ia import preparar_vistas_ia, preparar_vistas_ia_postgres

logger = obtener_logger("GestorCasos")

class GestorCasos:
    """
    Repositorio CRUD de datos para la lectura, actualización y persistencia del reporte.
    """
    COLUMNAS_MOLDE_EXPORTACION = [
        'FECHA INICIO JUICIO',
        'FECHA FIN ULTIMA FASE',
        'eta_id ULTIMA ETAPA',
        'ULTIMA ETAPA',
        'fas_id ULTIMA FASE',
        'ULTIMA FASE',
        'FECHA INICIO FASE ACTUAL',
        'eta_id ETAPA ACTUAL',
        'ETAPA ACTUAL',
        'fas_id FASE ACTUAL',
        'FASE ACTUAL',
        'DIAS TRANSCURRIDOS',
        CAMPO_FECHA_ULTIMA_GESTION,
        CAMPO_ESTADO_ULTIMA_GESTION,
    ]
    COLUMNAS_FECHA_PROCESAL = {
        'FECHA FIN ULTIMA FASE',
        'FECHA INICIO FASE ACTUAL',
        CAMPO_FECHA_ULTIMA_GESTION,
    }
    FILA_ENCABEZADO_REPORTE = 6

    @staticmethod
    def _parsear_fecha_reporte(valor):
        """Interpreta fechas del portal, incluidas marcas ISO con zona horaria."""
        if valor is None or pd.isna(valor) or str(valor).strip() == "":
            return None
        if isinstance(valor, pd.Timestamp):
            return valor.to_pydatetime()
        if isinstance(valor, datetime):
            return valor

        fecha_str = str(valor).strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(fecha_str[:10], fmt)
            except ValueError:
                continue
        try:
            return datetime.fromisoformat(fecha_str.replace("Z", "+00:00"))
        except ValueError:
            return None

    @classmethod
    def _normalizar_fecha_reporte(cls, valor):
        fecha = cls._parsear_fecha_reporte(valor)
        return fecha.strftime("%d/%m/%Y") if fecha else valor

    def __init__(self, ruta_config="config.json", *, config_efectiva=None):
        self.ruta_config = ruta_config
        if config_efectiva is None:
            with open(ruta_config, 'r', encoding='utf-8') as f:
                self.config = json.load(f)
        else:
            self.config = config_efectiva

        rutas = self.config.get('rutas', {})
        self.ruta_csv = rutas.get('archivo_csv', 'data/reporte_trabajo.csv')
        self.ruta_excel = rutas.get('archivo_origen', 'data/REPORTE JUICIOS PARA REVISIÓN JULIO.xlsx')
        self.ruta_final = rutas.get('archivo_excel_final', 'data/REPORTE_PROCESADO_FINAL.xlsx')
        self.hoja = rutas.get('hoja_lectura', 'migrado')
        self.filtros = self.config.get('filtros_activos', {})

        # Cargar CSV existente si existe y es válido
        if os.path.exists(self.ruta_csv):
            try:
                self.df = pd.read_csv(self.ruta_csv, low_memory=False)
                self.df.columns = self.df.columns.astype(str).str.strip().str.upper()
                if self.df.empty:
                    raise EmptyDataError("CSV vacío")
            except Exception as e:
                logger.warning("No se pudo cargar el CSV existente ('%s'): %s. Regenerando desde Excel...", self.ruta_csv, e)
                self._inicializar_csv(forzar=True)
                self.df = pd.read_csv(self.ruta_csv, low_memory=False)
                self.df.columns = self.df.columns.astype(str).str.strip().str.upper()
        else:
            logger.info("CSV no encontrado. Creando desde Excel original...")
            self._inicializar_csv(forzar=True)
            self.df = pd.read_csv(self.ruta_csv, low_memory=False)
            self.df.columns = self.df.columns.astype(str).str.strip().str.upper()

        # Comparar contra el total del reporte; otros archivos validos pueden
        # contener menos de 1000 filas.
        total_esperado = self.config.get('auditoria', {}).get('total_esperado')
        csv_incompleto = (
            total_esperado is not None and len(self.df) < int(total_esperado)
        )
        if ('SUCURSAL' not in self.df.columns or csv_incompleto) and os.path.exists(self.ruta_excel):
            logger.info("El CSV tiene %s registros. Sincronizando datos con el Excel completo...", len(self.df))
            self._inicializar_csv(forzar=True)
            self.df = pd.read_csv(self.ruta_csv, low_memory=False)
            self.df.columns = self.df.columns.astype(str).str.strip().str.upper()

        self.df = self._unificar_columnas_id(self.df)

    @staticmethod
    def _unificar_columnas_id(df):
        """Conserva un ID por campo; prioriza el valor no vacio mas reciente."""
        for destino in CAMPOS_ID:
            posiciones = []
            for indice, nombre in enumerate(df.columns):
                nombre = str(nombre).strip().upper()
                partes = nombre.rsplit('.', 1)
                if len(partes) == 2 and partes[1].isdigit():
                    nombre = partes[0]
                if nombre == destino.upper():
                    posiciones.append(indice)
            if not posiciones:
                continue
            valor = df.iloc[:, posiciones[0]].copy()
            for indice in posiciones[1:]:
                candidato = df.iloc[:, indice]
                valido = candidato.notna() & candidato.astype(str).str.strip().ne('')
                valor = candidato.where(valido, valor)
            mantener = [i for i in range(len(df.columns)) if i not in posiciones]
            df = df.iloc[:, mantener].copy()
            df[destino] = valor
        return df

    def _cargar_excel_robusto(self):
        """Carga el Excel usando una copia sombra para evitar bloqueos si está abierto en Excel, e infiere el header."""
        excel_path = os.path.abspath(self.ruta_excel)
        temp_path = None
        archivo_lectura = self.ruta_excel
        try:
            descriptor, temp_path = tempfile.mkstemp(
                prefix="_excel_shadow_",
                suffix=".xlsx",
                dir=os.path.dirname(excel_path) or None,
            )
            os.close(descriptor)
            shutil.copy2(excel_path, temp_path)
            archivo_lectura = temp_path
        except OSError as error:
            logger.warning(
                "No se pudo crear la copia sombra de Excel; se leerá el original: %s",
                error,
            )
            if temp_path and os.path.exists(temp_path):
                os.remove(temp_path)
            temp_path = None

        try:
            for h in [0, 1]:
                try:
                    df = pd.read_excel(archivo_lectura, sheet_name=self.hoja, header=h)
                    df.columns = df.columns.astype(str).str.strip().str.upper()
                    if 'SUCURSAL' in df.columns or 'CODIGO_JUICIO' in df.columns:
                        logger.info("Excel cargado correctamente detectando header=%s (%s filas).", h, len(df))
                        return df
                except Exception:
                    continue
            raise ValueError("No se pudo detectar la cabecera correcta en el archivo Excel.")
        finally:
            if temp_path and os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except Exception:
                    pass

    def _inicializar_csv(self, forzar=False):
        """CREATE: Genera o combina el CSV de trabajo desde el Excel original."""
        if not os.path.exists(self.ruta_excel):
            return

        if os.path.exists(self.ruta_csv) and not forzar:
            return

        logger.info("Inicializando/sincronizando base de datos CSV desde Excel...")
        df_excel = self._cargar_excel_robusto()

        if os.path.exists(self.ruta_csv):
            try:
                df_existente = pd.read_csv(self.ruta_csv, low_memory=False)
                df_existente.columns = df_existente.columns.astype(str).str.strip().str.upper()
                if 'CODIGO_JUICIO' in df_existente.columns and 'CODIGO_JUICIO' in df_excel.columns:
                    logger.info("Combinando datos existentes del CSV con la base completa del Excel...")
                    df_merged = df_existente.set_index('CODIGO_JUICIO').combine_first(df_excel.set_index('CODIGO_JUICIO')).reset_index()
                    df_merged.to_csv(self.ruta_csv, index=False, encoding='utf-8-sig')
                    return
            except Exception as e:
                logger.warning("No se pudo combinar el CSV existente, se creará uno nuevo: %s", e)

        df_excel.to_csv(self.ruta_csv, index=False, encoding='utf-8-sig')

    def _causas_hoja_seleccion(self):
        hoja = str(self.filtros.get('hoja_seleccion_causas') or '').strip()
        if not hoja:
            return None
        seleccion = pd.read_excel(self.ruta_excel, sheet_name=hoja, dtype=str)
        seleccion.columns = seleccion.columns.astype(str).str.strip().str.upper()
        if 'NUMERO_JUICIO' not in seleccion.columns:
            raise ValueError('HOJA_SELECCION_SIN_NUMERO_JUICIO')
        objetivos = {
            numero.replace('-', '').strip()
            for numero in seleccion['NUMERO_JUICIO'].dropna().astype(str)
            if numero.strip()
        }
        if not objetivos:
            raise ValueError('HOJA_SELECCION_SIN_CAUSAS')
        return objetivos

    def obtener_casos_pendientes(self):
        """READ: Obtiene la lista de números de juicio que cumplen con los filtros."""
        logger.debug("Columnas disponibles: %s", self.df.columns.tolist())
        columna_configurada = str(self.filtros.get('columna_estado_judicial', '')).strip().upper()
        if columna_configurada:
            if columna_configurada not in self.df.columns:
                raise KeyError(
                    f"La columna configurada para estado judicial '{columna_configurada}' no existe en el CSV."
                )
            col_estado = columna_configurada
            logger.info("Usando columna configurada para estado judicial: '%s'.", col_estado)
        elif 'ESTADO' in self.df.columns:
            col_estado = 'ESTADO'
        elif 'ESTADO.1' in self.df.columns:
            col_estado = 'ESTADO.1'
            logger.warning("Se usará la columna de respaldo 'ESTADO.1'.")
        else:
            raise KeyError("No se encontro una columna 'ESTADO' ni 'ESTADO.1' en el CSV.")
        
        mask = pd.Series(True, index=self.df.index)

        suc = str(self.filtros.get('sucursal', '') or '').strip().upper()
        if suc and suc not in ('TODAS', 'TODOS', 'ALL', 'NONE'):
            mask &= (self.df['SUCURSAL'].astype(str).str.strip().str.upper() == suc)

        ofi = str(self.filtros.get('oficina', '') or '').strip().upper()
        if ofi and ofi not in ('TODAS', 'TODOS', 'ALL', 'NONE'):
            mask &= (self.df['OFICINA'].astype(str).str.strip().str.upper() == ofi)
        usuario = str(self.filtros.get('usuario', '') or '').strip().upper()

        if usuario and usuario not in ('TODAS', 'TODOS', 'ALL', 'NONE'):
            if 'USUARIO' not in self.df.columns:
                raise KeyError("La columna configurada para usuario 'USUARIO' no existe en el CSV.")
            mask &= (
                self.df['USUARIO'].astype(str).str.strip().str.upper() == usuario
            )

        est = str(self.filtros.get('estado_judicial', '') or '').strip().upper()
        if est and est not in ('TODAS', 'TODOS', 'ALL', 'NONE'):
            mask &= (self.df[col_estado].astype(str).str.strip().str.upper() == est)

        df_filtrado = self.df[mask]

        casos = df_filtrado['NUMERO_JUICIO'].dropna().astype(str).str.strip().tolist()

        objetivos = self._causas_hoja_seleccion()
        if objetivos is not None:
            disponibles = {caso.replace('-', '').strip() for caso in casos}
            if not objetivos.issubset(disponibles):
                raise ValueError(
                    'CAUSAS_SELECCION_NO_DISPONIBLES:%s'
                    % len(objetivos - disponibles)
                )
            casos = [
                caso for caso in casos
                if caso.replace('-', '').strip() in objetivos
            ]

        revisiones = self.filtros.get('causas_revision_manual') or {}
        if not isinstance(revisiones, dict):
            raise ValueError('CAUSAS_REVISION_MANUAL_DEBE_SER_OBJETO')
        excluidas = {str(c).replace('-', '').strip() for c in revisiones}
        casos = [c for c in casos if c.replace('-', '').strip() not in excluidas]
        if objetivos is not None and len({c.replace('-', '').strip() for c in casos}) != len(objetivos):
            raise ValueError('CAUSAS_SELECCION_EXCLUIDAS_REVISION_MANUAL')

        # Aplicar punto de partida si fue especificado
        inicio = self.filtros.get('inicio_desde_juicio')
        if inicio:
            inicio_limpio = str(inicio).replace("-", "").strip()
            idx = next((i for i, c in enumerate(casos) if str(c).replace("-", "").strip() == inicio_limpio), None)
            if idx is not None:
                logger.info("Reanudando desde causa '%s' (Caso #%s de %s).", inicio, idx + 1, len(casos))
                casos = casos[idx:]

        return casos

    def actualizar_caso(self, numero_juicio, datos):
        """UPDATE: Inyecta en la fila correspondiente los datos extraídos."""
        datos = dict(datos)
        enriquecer_datos_procesales(datos)
        enriquecer_ultima_gestion_judicial(datos)
        numero_juicio_normalizado = str(numero_juicio).strip().upper()
        mask = (
            self.df['NUMERO_JUICIO'].astype(str).str.strip().str.upper()
            == numero_juicio_normalizado
        )
        if mask.any():
            for col, val in datos.items():
                if val is not None:
                    if col in self.COLUMNAS_FECHA_PROCESAL:
                        val = self._normalizar_fecha_reporte(val)
                    if col not in self.df.columns:
                        self.df[col] = None

                    # Si el valor es una lista o diccionario (ej. HISTORIAL_ACTUACIONES), serializar a JSON string
                    if isinstance(val, (list, dict)):
                        val = json.dumps(val, ensure_ascii=False)
                    # Un CSV con una columna completamente vacia se carga como
                    # float64. Pandas recientes rechazan guardar texto (por
                    # ejemplo, una fecha dd/mm/aaaa) sin convertir antes el tipo.
                    if isinstance(val, str) and self.df[col].dtype != object:
                        self.df[col] = self.df[col].astype(object)


                    self.df.loc[mask, col] = val
            return True
        return False

    def depurar_filas_duplicadas_exactas(self):
        """Elimina solo registros iguales en todas sus columnas.

        Un mismo nÃºmero de juicio puede pertenecer a mÃ¡s de una cartera,
        usuario o crÃ©dito. Por eso las coincidencias por ``NUMERO_JUICIO`` no
        se eliminan automÃ¡ticamente: Ãºnicamente se depuran copias exactas,
        tratando vacÃ­os y nulos como equivalentes.
        """
        if self.df.empty:
            return 0

        comparable = self.df.fillna("").astype(str)
        duplicadas = comparable.duplicated(keep="first")
        cantidad = int(duplicadas.sum())
        if cantidad:
            self.df = self.df.loc[~duplicadas].copy()
            logger.info("Se depuraron %s filas duplicadas exactas.", cantidad)
        return cantidad


    @staticmethod
    def _es_bloqueo_transitorio_archivo(error):
        """Reconoce bloqueos habituales de Windows, OneDrive y antivirus."""
        return (
            getattr(error, "winerror", None) in {32, 33}
            or getattr(error, "errno", None) in {errno.EACCES, errno.EBUSY}
        )

    def _crear_respaldo_csv(self):
        """Crea el .bak y tolera bloqueos breves sin omitir la salvaguarda."""
        ruta_backup = f"{self.ruta_csv}.bak"
        max_intentos = 6
        for intento in range(1, max_intentos + 1):
            try:
                shutil.copy2(self.ruta_csv, ruta_backup)
                logger.debug("Respaldo del CSV creado en: %s", ruta_backup)
                return True
            except OSError as error:
                es_transitorio = self._es_bloqueo_transitorio_archivo(error)
                es_ultimo_intento = intento == max_intentos
                if not es_transitorio or es_ultimo_intento:
                    logger.error("No se pudo crear respaldo del CSV: %s", error)
                    return False
                espera = min(0.5 * (2 ** (intento - 1)), 2.0)
                logger.warning(
                    "[RESPALDO_CSV_REINTENTO] Archivo temporalmente bloqueado; "
                    "reintentando en %.1fs (%s/%s): %s",
                    espera, intento, max_intentos, error,
                )
                time.sleep(espera)
        return False

    def guardar(self):
        """SAVE: Reemplaza atómicamente el CSV tras crear su respaldo."""
        if os.path.isfile(self.ruta_csv) and not self._crear_respaldo_csv():
            return False

        ruta_temporal = None
        try:
            ruta_absoluta = os.path.abspath(self.ruta_csv)
            descriptor, ruta_temporal = tempfile.mkstemp(
                prefix=".%s." % os.path.basename(ruta_absoluta),
                suffix=".tmp",
                dir=os.path.dirname(ruta_absoluta) or None,
            )
            os.close(descriptor)
            self.df.to_csv(
                ruta_temporal,
                index=False,
                encoding='utf-8-sig',
            )
            os.replace(ruta_temporal, ruta_absoluta)
            ruta_temporal = None
            return True
        except Exception as e:
            logger.error("No se pudo guardar CSV: %s", e)
            return False
        finally:
            if ruta_temporal and os.path.exists(ruta_temporal):
                try:
                    os.remove(ruta_temporal)
                except OSError:
                    logger.warning(
                        "No se pudo eliminar el CSV temporal: %s",
                        ruta_temporal,
                    )

    @staticmethod
    def _mascara_mal_ingresado(df):
        comentarios = df.get(
            'COMENTARIO_ULTIMO', pd.Series('', index=df.index)
        ).fillna('').astype(str).str.upper()
        mascara = (
            comentarios.str.contains('FORMATO_CAUSA_INVALIDO', regex=False)
            | (
                comentarios.str.contains('NO DEVOLVI', regex=False)
                & comentarios.str.contains('RESULTAD', regex=False)
            )
        )
        for campo_estado in ('ETAPA ACTUAL', 'FASE ACTUAL'):
            if campo_estado in df.columns:
                mascara |= (
                    df[campo_estado].fillna('').astype(str).str.strip().str.upper()
                    == 'EXCLUIDO_NO_CORRESPONDE'
                )
        for campo in (
            'FECHA INICIO JUICIO', 'FECHA FIN ULTIMA FASE', 'ULTIMA ETAPA',
            'ULTIMA FASE', 'FECHA INICIO FASE ACTUAL', 'ETAPA ACTUAL',
            'FASE ACTUAL',
        ):
            if campo in df.columns:
                mascara |= (
                    df[campo].fillna('').astype(str).str.strip().str.upper()
                    == 'MAL INGRESADO'
                )
        return mascara

    @staticmethod
    def _agregar_motivo_id_no_catalogado(df, mascara_mal_ingresado):
        pares = (
            ('ULTIMA ETAPA', 'eta_id ULTIMA ETAPA'),
            ('ULTIMA FASE', 'fas_id ULTIMA FASE'),
            ('ETAPA ACTUAL', 'eta_id ETAPA ACTUAL'),
            ('FASE ACTUAL', 'fas_id FASE ACTUAL'),
        )
        if 'COMENTARIO_ULTIMO' not in df.columns:
            df['COMENTARIO_ULTIMO'] = None
        for idx, fila in df.loc[~mascara_mal_ingresado].iterrows():
            faltantes = []
            for etiqueta, campo_id in pares:
                valor = fila.get(etiqueta)
                if valor is not None and not pd.isna(valor) and str(valor).strip():
                    if fila.get(campo_id) is None or pd.isna(fila.get(campo_id)):
                        faltantes.append(f"{etiqueta}='{valor}'")
            if not faltantes:
                continue
            motivo = 'ID_NO_CATALOGADO: ' + ', '.join(faltantes)
            comentario = fila.get('COMENTARIO_ULTIMO')
            comentario = '' if comentario is None or pd.isna(comentario) else str(comentario).strip()
            if motivo not in comentario:
                df.at[idx, 'COMENTARIO_ULTIMO'] = (
                    f"{comentario} | {motivo}" if comentario else motivo
                )

    def _preparar_exportacion(self, fecha_actual=None):
        self.calcular_dias_transcurridos(fecha_actual)
        df_export = self.df.copy()
        if getattr(self, 'filtros', None):
            objetivos = self._causas_hoja_seleccion()
            if objetivos is not None:
                numeros = df_export['NUMERO_JUICIO'].astype(str).str.replace(
                    '-', '', regex=False,
                ).str.strip()
                df_export = df_export.loc[numeros.isin(objetivos)].copy()
                encontrados = set(numeros.loc[df_export.index])
                if encontrados != objetivos or len(df_export) != len(objetivos):
                    raise ValueError('EXPORTACION_SELECCION_INCOMPLETA_O_DUPLICADA')
        nuevas_cols = list(self.COLUMNAS_MOLDE_EXPORTACION)
        for columna in nuevas_cols:
            if columna not in df_export.columns:
                df_export[columna] = None
        for columna in (
            'COMENTARIO_ULTIMO', 'FECHA INICIO JUICIO',
            'FECHA FIN ULTIMA FASE', 'ULTIMA ETAPA', 'ULTIMA FASE',
            'FECHA INICIO FASE ACTUAL', 'ETAPA ACTUAL', 'FASE ACTUAL',
        ):
            if columna in df_export.columns and df_export[columna].dtype != object:
                df_export[columna] = df_export[columna].astype(object)

        for idx, fila in df_export.iterrows():
            datos = fila.to_dict()
            enriquecer_datos_procesales(datos, normalizar_etiquetas=True)
            enriquecer_ultima_gestion_judicial(datos)
            for campo in (
                'ULTIMA ETAPA', 'ULTIMA FASE', 'ETAPA ACTUAL', 'FASE ACTUAL',
                CAMPO_FECHA_ULTIMA_GESTION,
                CAMPO_ESTADO_ULTIMA_GESTION,
                *CAMPOS_ID,
            ):
                df_export.at[idx, campo] = datos.get(campo)

        for col_fecha in self.COLUMNAS_FECHA_PROCESAL:
            df_export[col_fecha] = df_export[col_fecha].map(
                self._normalizar_fecha_reporte
            )
        mask_copia = (
            df_export['FECHA INICIO FASE ACTUAL'].isna()
            & df_export['FECHA FIN ULTIMA FASE'].notna()
        )
        df_export.loc[mask_copia, 'FECHA INICIO FASE ACTUAL'] = df_export.loc[
            mask_copia, 'FECHA FIN ULTIMA FASE'
        ]

        columnas_obsoletas = [
            'ETAPA_PROCESAL (ACTUAL)', 'FASE_PROCESAL (ACTUAL)',
            'ETAPA_PROCESAL (MIGRADO)', 'CODIGO_FASE',
            'FASE_PROCESAL (MIGRADO)', 'FECHA INICIAL FASE ACTUAL',
            'DIAS EN LA FASE ACTUAL', 'ETAPA_PROCESAL', 'FASE_PROCESAL', ' ',
            'FECHA_ULTIMA_GESTION_JUDICIAL',
        ]
        df_export.drop(columns=columnas_obsoletas, inplace=True, errors='ignore')

        if 'CODIGO_JUICIO' in df_export.columns:
            codigos = df_export['CODIGO_JUICIO']
            if codigos.isna().any() or codigos.astype(str).str.strip().eq('').any():
                raise ValueError('CODIGO_JUICIO_VACIO_EN_EXPORTACION')
            if codigos.astype(str).str.strip().duplicated().any():
                raise ValueError('CODIGO_JUICIO_DUPLICADO_EN_EXPORTACION')
        else:
            logger.warning("El reporte no contiene CODIGO_JUICIO; no se pudo validar su identidad.")

        mascara_mal = self._mascara_mal_ingresado(df_export)
        campos_mal = [
            'FECHA INICIO JUICIO', 'FECHA FIN ULTIMA FASE', 'ULTIMA ETAPA',
            'ULTIMA FASE', 'FECHA INICIO FASE ACTUAL', 'ETAPA ACTUAL',
            'FASE ACTUAL',
        ]
        df_export.loc[mascara_mal, campos_mal] = 'MAL INGRESADO'
        df_export.loc[mascara_mal, [*CAMPOS_ID, 'DIAS TRANSCURRIDOS']] = None
        mascara_revision_titulo = df_export[CAMPO_ESTADO_ULTIMA_GESTION].eq(
            MARCA_REVISION_ULTIMA_GESTION
        )
        for idx in df_export.index[mascara_revision_titulo]:
            valor_comentario = df_export.at[idx, 'COMENTARIO_ULTIMO']
            comentario = (
                '' if pd.isna(valor_comentario) else str(valor_comentario).strip()
            )
            if MOTIVO_REVISION_ULTIMA_GESTION not in comentario:
                marca = f'REVISION MANUAL: {MOTIVO_REVISION_ULTIMA_GESTION}'
                df_export.at[idx, 'COMENTARIO_ULTIMO'] = (
                    f'{comentario} | {marca}' if comentario else marca
                )
        self._agregar_motivo_id_no_catalogado(df_export, mascara_mal)

        cols_base = [c for c in df_export.columns if c not in nuevas_cols]
        cols_finales = cols_base + nuevas_cols
        reporte = df_export[cols_finales]
        para_carga = reporte.loc[~(mascara_mal | mascara_revision_titulo)].copy()
        return reporte, para_carga

    @staticmethod
    def _formatear_errores_excel(ruta, generado_en=None):
        """Aplica el acabado profesional del libro sin modificar sus datos."""
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter

        wb = openpyxl.load_workbook(ruta)
        generado_en = generado_en or datetime.now()

        colores = {
            'azul': 'FF17365D',
            'pizarra': 'FF34495E',
            'dorado': 'FFC9A227',
            'azul_claro': 'FFD9E2F3',
            'gris': 'FFF3F5F7',
            'borde': 'FFD9E0E7',
            'rojo_fondo': 'FFFCE8E6',
            'rojo_texto': 'FF9C0006',
            'ambar_fondo': 'FFFFF4CC',
            'ambar_texto': 'FF7F6000',
            'verde_fondo': 'FFE8F3EC',
            'verde_texto': 'FF1F6D42',
            'blanco': 'FFFFFFFF',
            'texto': 'FF1F2933',
        }
        fill = {
            clave: PatternFill('solid', fgColor=valor)
            for clave, valor in colores.items()
            if clave not in {'borde', 'texto', 'rojo_texto', 'ambar_texto', 'verde_texto'}
        }
        borde_suave = Side(style='thin', color=colores['borde'])
        borde_dorado = Side(style='medium', color=colores['dorado'])

        anchos = {
            'CODIGO_JUICIO': 15, 'SUCURSAL': 14, 'OFICINA': 16,
            'ASESOR': 16, 'USUARIO': 14, 'CREDITO': 18,
            'CEDULA_IDENTIDAD': 17, 'ESTADO': 13, 'SEGMENTO': 15,
            'NOMBRES_CLIENTE': 28, 'NOMBRE': 28, 'SALDO_CAPITAL': 16,
            'SALDO_TOTAL': 16, 'SALDO_CLIENTE_GRUPO': 20,
            'PRODUCTO': 18, 'NUMERO_JUICIO': 24, 'JUZGADO': 32,
            'CUANTIA': 16, 'DIAS_MORA': 12, 'CODIGO_ETAPA': 15,
            'FECHA_INICIO': 17, 'FECHA_ULTIMA_GESTION_JUDICIAL': 24,
            'FECHA_ULTIMO_COMPROMISO': 22, 'VALOR_ULTIMO_COMPROMISO': 22,
            'RECUPERACION_ACTUAL': 20, 'RECUPERACION_MES_ANTERIOR': 24,
            'SISTEMA': 14, 'COMENTARIO_ULTIMO': 38,
            'HISTORIAL_ACTUACIONES': 45, 'FECHA INICIO JUICIO': 19,
            'FECHA FIN ULTIMA FASE': 21, 'eta_id ULTIMA ETAPA': 19,
            'ULTIMA ETAPA': 30, 'fas_id ULTIMA FASE': 18,
            'ULTIMA FASE': 32, 'FECHA INICIO FASE ACTUAL': 23,
            'eta_id ETAPA ACTUAL': 19, 'ETAPA ACTUAL': 30,
            'fas_id FASE ACTUAL': 18, 'FASE ACTUAL': 32,
            'DIAS TRANSCURRIDOS': 19,
            CAMPO_FECHA_ULTIMA_GESTION: 24,
            CAMPO_ESTADO_ULTIMA_GESTION: 38,
            'ESTADO AUDITORIA IA': 22, 'FUENTE DECISION': 24,
            'CONFIANZA IA': 17, 'REVISION PENDIENTE': 22,
            'NUMERO_JUICIO': 24, 'ACTUACION_ID': 31, 'CARPETA': 22,
            'CONDICION': 18, 'FUENTE': 18, 'VERSION': 22,
            'EVIDENCIAS': 42, 'MOTIVO IA': 55,
            'OBSERVACION HUMANA': 45,
        }
        columnas_moneda = {
            'SALDO_CAPITAL', 'SALDO_TOTAL', 'SALDO_CLIENTE_GRUPO', 'CUANTIA',
            'VALOR_ULTIMO_COMPROMISO', 'RECUPERACION_ACTUAL',
            'RECUPERACION_MES_ANTERIOR',
        }
        columnas_identificador = {
            'CODIGO_JUICIO', 'CREDITO', 'CEDULA_IDENTIDAD', 'NUMERO_JUICIO',
        }
        columnas_enteras = {
            'DIAS_MORA', 'CODIGO_ETAPA', 'eta_id ULTIMA ETAPA',
            'fas_id ULTIMA FASE', 'eta_id ETAPA ACTUAL',
            'fas_id FASE ACTUAL', 'DIAS TRANSCURRIDOS',
        }
        columnas_fecha = {
            'FECHA_INICIO', 'FECHA_ULTIMA_GESTION_JUDICIAL',
            'FECHA_ULTIMO_COMPROMISO', 'FECHA INICIO JUICIO',
            'FECHA FIN ULTIMA FASE', 'FECHA INICIO FASE ACTUAL',
            CAMPO_FECHA_ULTIMA_GESTION,
        }
        columnas_envueltas = {
            'NOMBRES_CLIENTE', 'NOMBRE', 'JUZGADO', 'COMENTARIO_ULTIMO',
            'HISTORIAL_ACTUACIONES', 'ULTIMA ETAPA', 'ULTIMA FASE',
            'ETAPA ACTUAL', 'FASE ACTUAL', CAMPO_ESTADO_ULTIMA_GESTION,
            'EVIDENCIAS', 'MOTIVO IA', 'OBSERVACION HUMANA',
        }

        def es_error_visual(comentario, valores_estado):
            return (
                'ERROR:' in comentario
                or 'ERROR_' in comentario
                or 'FORMATO_CAUSA_INVALIDO' in comentario
                or ('NO DEVOLVI' in comentario and 'RESULTAD' in comentario)
                or 'MAL INGRESADO' in valores_estado
                or 'EXCLUIDO_NO_CORRESPONDE' in valores_estado
            )

        def aplicar_marco_kpi(ws, columnas, etiqueta, valor, color_valor):
            inicio, fin = columnas
            rango_etiqueta = f'{inicio}3:{fin}3'
            rango_valor = f'{inicio}4:{fin}4'
            ws.merge_cells(rango_etiqueta)
            ws.merge_cells(rango_valor)
            ws[f'{inicio}3'] = etiqueta
            ws[f'{inicio}4'] = valor
            for fila in (3, 4):
                for row in ws.iter_rows(
                    min_row=fila, max_row=fila,
                    min_col=ws[f'{inicio}{fila}'].column,
                    max_col=ws[f'{fin}{fila}'].column,
                ):
                    for celda in row:
                        celda.fill = fill['azul_claro'] if fila == 3 else PatternFill(
                            'solid', fgColor=color_valor
                        )
                        celda.border = Border(
                            left=borde_suave, right=borde_suave,
                            top=borde_suave, bottom=borde_suave,
                        )
            ws[f'{inicio}3'].font = Font(
                name='Arial', size=9, bold=True, color=colores['pizarra']
            )
            ws[f'{inicio}4'].font = Font(
                name='Arial', size=14, bold=True, color=colores['texto']
            )
            ws[f'{inicio}3'].alignment = Alignment(horizontal='center', vertical='center')
            ws[f'{inicio}4'].alignment = Alignment(horizontal='center', vertical='center')

        def estilizar_hoja(ws, fila_encabezado, nombre_tabla, color_pestana):
            encabezados = {
                str(celda.value).strip(): celda.column
                for celda in ws[fila_encabezado] if celda.value is not None
            }
            if not encabezados:
                return

            ultima_columna = max(encabezados.values())
            ultima_fila = ws.max_row
            letra_final = get_column_letter(ultima_columna)
            primera_fila_datos = fila_encabezado + 1
            inicio_procesal = encabezados.get('FECHA INICIO JUICIO', ultima_columna + 1)

            ws.sheet_view.showGridLines = False
            ws.sheet_view.zoomScale = 85
            ws.sheet_properties.tabColor = color_pestana
            ws.freeze_panes = f'A{primera_fila_datos}'
            ws.sheet_format.defaultRowHeight = 18
            ws.row_dimensions[fila_encabezado].height = 38

            for col_idx in range(1, ultima_columna + 1):
                celda = ws.cell(fila_encabezado, col_idx)
                celda.fill = fill['azul'] if col_idx < inicio_procesal else fill['pizarra']
                celda.font = Font(
                    name='Arial', size=9, bold=True, color=colores['blanco']
                )
                celda.alignment = Alignment(
                    horizontal='center', vertical='center', wrap_text=True
                )
                celda.border = Border(
                    left=Side(style='thin', color=colores['blanco']),
                    right=Side(style='thin', color=colores['blanco']),
                    bottom=borde_dorado,
                )
                encabezado = str(celda.value).strip()
                ancho = anchos.get(encabezado, min(max(len(encabezado) + 2, 12), 22))
                ws.column_dimensions[get_column_letter(col_idx)].width = ancho

            col_comentario = encabezados.get('COMENTARIO_ULTIMO')
            cols_estado = [
                encabezados[c] for c in (
                    'ETAPA ACTUAL', 'FASE ACTUAL', 'DIAS TRANSCURRIDOS'
                ) if c in encabezados
            ]
            for row_idx in range(primera_fila_datos, ultima_fila + 1):
                comentario = str(
                    ws.cell(row_idx, col_comentario).value or ''
                ).upper() if col_comentario else ''
                valores_estado = ' '.join(
                    str(ws.cell(row_idx, col).value or '').upper()
                    for col in cols_estado
                )
                es_error = es_error_visual(comentario, valores_estado)
                es_revision = any(marca in comentario for marca in (
                    'REVISION MANUAL', 'ID_NO_CATALOGADO',
                ))
                contenido_largo = False

                for col_idx in range(1, ultima_columna + 1):
                    celda = ws.cell(row_idx, col_idx)
                    encabezado = str(ws.cell(fila_encabezado, col_idx).value).strip()
                    celda.font = Font(name='Arial', size=9.5, color=colores['texto'])
                    celda.alignment = Alignment(
                        horizontal='left', vertical='center',
                        wrap_text=encabezado in columnas_envueltas,
                    )
                    if isinstance(celda.value, (int, float)) and encabezado not in columnas_identificador:
                        celda.alignment = Alignment(horizontal='right', vertical='center')
                    if encabezado in columnas_moneda:
                        celda.number_format = '"$"#,##0.00'
                    elif encabezado in columnas_identificador:
                        celda.number_format = '0' if isinstance(celda.value, (int, float)) else '@'
                    elif encabezado in columnas_enteras:
                        celda.number_format = '#,##0'
                    elif encabezado in columnas_fecha and isinstance(
                        celda.value, (datetime, pd.Timestamp)
                    ):
                        celda.number_format = 'dd/mm/yyyy'
                    if encabezado in columnas_envueltas and len(str(celda.value or '')) > 80:
                        contenido_largo = True

                if es_error:
                    for col_idx in range(1, ultima_columna + 1):
                        celda = ws.cell(row_idx, col_idx)
                        celda.fill = fill['rojo_fondo']
                        celda.font = Font(
                            name='Arial', size=9.5, color=colores['rojo_texto'],
                            bold=True,
                        )
                elif es_revision:
                    columnas_alerta = set(cols_estado)
                    if col_comentario:
                        columnas_alerta.add(col_comentario)
                    for col_idx in columnas_alerta:
                        celda = ws.cell(row_idx, col_idx)
                        celda.fill = fill['ambar_fondo']
                        celda.font = Font(
                            name='Arial', size=9.5, color=colores['ambar_texto'],
                            bold=True,
                        )
                else:
                    for col_idx in cols_estado:
                        celda = ws.cell(row_idx, col_idx)
                        celda.fill = fill['verde_fondo']
                        celda.font = Font(
                            name='Arial', size=9.5, color=colores['verde_texto'],
                            bold=True,
                        )
                if contenido_largo:
                    ws.row_dimensions[row_idx].height = 32

            referencia = f'A{fila_encabezado}:{letra_final}{ultima_fila}'
            tabla = Table(displayName=nombre_tabla, ref=referencia)
            tabla.tableStyleInfo = TableStyleInfo(
                name='TableStyleMedium2', showFirstColumn=False,
                showLastColumn=False, showRowStripes=True,
                showColumnStripes=False,
            )
            ws.add_table(tabla)

            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
            ws.page_setup.fitToWidth = 2
            ws.page_setup.fitToHeight = 0
            ws.page_margins.left = 0.25
            ws.page_margins.right = 0.25
            ws.page_margins.top = 0.45
            ws.page_margins.bottom = 0.45
            ws.page_margins.header = 0.2
            ws.page_margins.footer = 0.2
            ws.print_title_rows = f'{fila_encabezado}:{fila_encabezado}'
            ws.print_area = f'A1:{letra_final}{ultima_fila}'
            ws.oddHeader.center.text = '&BReporte de gestión judicial'
            ws.oddFooter.left.text = 'Sistema Judicial'
            ws.oddFooter.center.text = 'Página &P de &N'
            ws.oddFooter.right.text = generado_en.strftime('%d/%m/%Y')

        ws_reporte = wb['Reporte']
        ws_carga = wb['PARA CARGA']
        fila_reporte = GestorCasos.FILA_ENCABEZADO_REPORTE

        encabezados_reporte = {
            str(celda.value).strip(): celda.column
            for celda in ws_reporte[fila_reporte] if celda.value is not None
        }
        total = max(ws_reporte.max_row - fila_reporte, 0)
        cargables = max(ws_carga.max_row - 1, 0)
        observados = max(total - cargables, 0)
        col_comentario = encabezados_reporte.get('COMENTARIO_ULTIMO')
        cols_estado_reporte = [
            encabezados_reporte[c] for c in ('ETAPA ACTUAL', 'FASE ACTUAL')
            if c in encabezados_reporte
        ]
        revision = 0
        if col_comentario:
            for row_idx in range(fila_reporte + 1, ws_reporte.max_row + 1):
                comentario = str(
                    ws_reporte.cell(row_idx, col_comentario).value or ''
                ).upper()
                valores_estado = ' '.join(
                    str(ws_reporte.cell(row_idx, col).value or '').upper()
                    for col in cols_estado_reporte
                )
                es_error = es_error_visual(comentario, valores_estado)
                if not es_error and any(marca in comentario for marca in (
                    'REVISION MANUAL', 'ID_NO_CATALOGADO',
                )):
                    revision += 1

        ultima_columna = max(encabezados_reporte.values())
        letra_final = get_column_letter(ultima_columna)
        ws_reporte.merge_cells(f'A1:{letra_final}1')
        ws_reporte.merge_cells(f'A2:{letra_final}2')
        ws_reporte['A1'] = 'REPORTE FINAL DE GESTIÓN JUDICIAL'
        ws_reporte['A2'] = (
            'Generado el ' + generado_en.strftime('%d/%m/%Y a las %H:%M')
        )
        for row in ws_reporte.iter_rows(min_row=1, max_row=2, min_col=1, max_col=ultima_columna):
            for celda in row:
                celda.fill = fill['azul']
        ws_reporte['A1'].font = Font(
            name='Arial', size=15, bold=True, color=colores['blanco']
        )
        ws_reporte['A2'].font = Font(
            name='Arial', size=9.5, italic=True, color=colores['blanco']
        )
        ws_reporte['A1'].alignment = Alignment(horizontal='left', vertical='center')
        ws_reporte['A2'].alignment = Alignment(horizontal='left', vertical='center')
        ws_reporte.row_dimensions[1].height = 28
        ws_reporte.row_dimensions[2].height = 20
        ws_reporte.row_dimensions[3].height = 18
        ws_reporte.row_dimensions[4].height = 25
        ws_reporte.row_dimensions[5].height = 8

        espacio_disponible = ultima_columna - 3
        tamano_base, sobrantes = divmod(espacio_disponible, 4)
        inicio = 1
        tarjetas = (
            ('Total de registros', total, colores['blanco']),
            ('Listos para carga', cargables, colores['verde_fondo']),
            ('Observados / no cargables', observados, colores['rojo_fondo']),
            ('Revisión manual / ID pendiente', revision, colores['ambar_fondo']),
        )
        for indice, (etiqueta, valor, color_valor) in enumerate(tarjetas):
            tamano = tamano_base + (1 if indice < sobrantes else 0)
            fin = inicio + tamano - 1
            aplicar_marco_kpi(
                ws_reporte,
                (get_column_letter(inicio), get_column_letter(fin)),
                etiqueta,
                valor,
                color_valor,
            )
            inicio = fin + 2

        estilizar_hoja(
            ws_reporte, fila_reporte, 'TablaReporteJudicial', colores['azul']
        )
        estilizar_hoja(
            ws_carga, 1, 'TablaParaCarga', colores['dorado']
        )
        if 'HISTORIAL HITOS' in wb:
            estilizar_hoja(wb['HISTORIAL HITOS'], 1, 'TablaHitosIA', colores['pizarra'])
        if 'REVISION IA' in wb:
            estilizar_hoja(wb['REVISION IA'], 1, 'TablaRevisionIA', colores['dorado'])
        wb.properties.title = 'Reporte final de gestión judicial'
        wb.properties.subject = 'Seguimiento procesal y archivo para carga a Sistemas'
        wb.properties.creator = 'Sistema Judicial ESPOIR'
        wb.save(ruta)
        wb.close()

    def exportar_excel(self, fecha_actual=None):
        """Genera el reporte, la hoja para Sistemas y las vistas de IA."""
        reporte, para_carga = self._preparar_exportacion(fecha_actual)
        hitos = revisiones = None
        configuracion = getattr(self, 'config', {})
        if configuracion.get('inferencia_ia'):
            try:
                if configuracion.get('base_de_datos', {}).get('motor') == 'postgres':
                    reporte, hitos, revisiones = preparar_vistas_ia_postgres(
                        configuracion, reporte
                    )
                else:
                    ruta_db = configuracion.get('rutas', {}).get('archivo_db')
                    if ruta_db:
                        if not os.path.isabs(ruta_db):
                            ruta_db = os.path.join(
                                os.path.dirname(os.path.abspath(getattr(self, 'ruta_config', 'config.json'))), ruta_db
                            )
                        reporte, hitos, revisiones = preparar_vistas_ia(ruta_db, reporte)
            except Exception as exc:
                logger.warning('No se pudieron preparar las vistas IA: %s', type(exc).__name__)
        ruta_objetivo = os.path.abspath(self.ruta_final)
        directorio = os.path.dirname(ruta_objetivo) or os.getcwd()
        os.makedirs(directorio, exist_ok=True)
        descriptor, ruta_temporal = tempfile.mkstemp(
            prefix=f".{os.path.basename(ruta_objetivo)}.",
            suffix='.xlsx',
            dir=directorio,
        )
        os.close(descriptor)
        try:
            with pd.ExcelWriter(ruta_temporal, engine='openpyxl') as writer:
                reporte.to_excel(
                    writer, index=False, sheet_name='Reporte',
                    startrow=self.FILA_ENCABEZADO_REPORTE - 1,
                )
                para_carga.to_excel(writer, index=False, sheet_name='PARA CARGA')
                if hitos is not None:
                    hitos.to_excel(writer, index=False, sheet_name='HISTORIAL HITOS')
                    revisiones.to_excel(writer, index=False, sheet_name='REVISION IA')
            self._formatear_errores_excel(ruta_temporal)
            try:
                os.replace(ruta_temporal, ruta_objetivo)
            except PermissionError:
                base, ext = os.path.splitext(ruta_objetivo)
                ruta_objetivo = f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
                logger.warning(
                    "Archivo Excel bloqueado. Guardando copia en: %s", ruta_objetivo
                )
                os.replace(ruta_temporal, ruta_objetivo)
            ruta_temporal = None
            self.ruta_final = ruta_objetivo
            logger.info(
                "Excel generado: %s filas en Reporte y %s en PARA CARGA.",
                len(reporte), len(para_carga),
            )
        finally:
            if ruta_temporal and os.path.exists(ruta_temporal):
                os.remove(ruta_temporal)

    def calcular_dias_transcurridos(self, fecha_actual=None):
        """
        Calcula la columna 'DIAS TRANSCURRIDOS' como la diferencia en días calendario
        entre la fecha actual y 'FECHA INICIO FASE ACTUAL'. Usa la fecha de
        fin de la última fase solo como respaldo.
        Soporta formatos dd/mm/yyyy y yyyy-mm-dd.
        """
        col_dias = 'DIAS TRANSCURRIDOS'
        columnas_fecha = [
            col for col in (
                'FECHA INICIO FASE ACTUAL', 'FECHA FIN ULTIMA FASE',
                'FECHA INICIAL FASE ACTUAL',
            ) if col in self.df.columns
        ]
        if not columnas_fecha:
            logger.warning("No hay fecha procesal para calcular '%s'.", col_dias)
            return

        self.df[col_dias] = None

        referencia = fecha_actual or datetime.now()
        hoy = referencia.date() if isinstance(referencia, datetime) else referencia
        conteo = 0

        for idx, fila in self.df.iterrows():
            valor = next(
                (
                    fila.get(columna) for columna in columnas_fecha
                    if fila.get(columna) is not None
                    and not pd.isna(fila.get(columna))
                    and str(fila.get(columna)).strip()
                ),
                None,
            )
            fecha_parsed = self._parsear_fecha_reporte(valor)

            if fecha_parsed:
                dias = (hoy - fecha_parsed.date()).days
                self.df.at[idx, col_dias] = max(0, dias)
                conteo += 1

        logger.info("'%s' calculado para %s registros.", col_dias, conteo)

    def calcular_dias_fase_actual(self):
        """Método de retrocompatibilidad que delega en calcular_dias_transcurridos."""
        self.calcular_dias_transcurridos()
